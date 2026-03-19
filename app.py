import io
import json
import matplotlib.pyplot as plt
import pandas as pd
import streamlit as st

from calculations import calculate_traverse
from sample_data import SAMPLE_PROJECT, SAMPLE_ROWS
from validation import validate_project_settings, validate_traverse_rows


def build_setup_rows(num_setups):
    rows = []
    for i in range(1, num_setups + 1):
        rows.append({
            "setup_number": i,
            "from_point": i,
            "to_point": i + 1 if i < num_setups else 1,
            "angle_deg": None,
            "angle_min": None,
            "angle_sec": None,
            "distance": None,
        })
    return pd.DataFrame(rows)


def clean_first_setup_angles(df):
    if not df.empty and "setup_number" in df.columns:
        for idx in df.index:
            try:
                if pd.notna(df.at[idx, "setup_number"]) and int(float(df.at[idx, "setup_number"])) == 1:
                    df.at[idx, "angle_deg"] = None
                    df.at[idx, "angle_min"] = None
                    df.at[idx, "angle_sec"] = None
            except Exception:
                pass
    return df


def df_to_csv_download(df):
    return df.to_csv(index=False).encode("utf-8")


def build_traverse_plot(coordinates_df, traverse_name="Traverse Plot"):
    fig, ax = plt.subplots(figsize=(10, 5))

    x = coordinates_df["easting"]
    y = coordinates_df["northing"]
    labels = coordinates_df["point"]

    ax.plot(x, y, marker="o", linestyle="-", linewidth=2, markersize=6)

    for xi, yi, label in zip(x, y, labels):
        ax.annotate(str(label), (xi, yi), textcoords="offset points", xytext=(5, 5))

    ax.set_title(traverse_name)
    ax.set_xlabel("Easting")
    ax.set_ylabel("Northing")
    ax.grid(True, linestyle="--", alpha=0.5)
    ax.set_aspect("equal", adjustable="datalim")

    plt.tight_layout()
    return fig

def build_excel_workbook(results):
    output = io.BytesIO()

    raw_df = results["raw_df"].copy()
    adjusted_df = results["adjusted_df"].copy()
    coordinates_df = results["coordinates_df"].copy()
    closure = results["closure"]

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        workbook = writer.book

        from openpyxl.chart import ScatterChart, Reference, Series
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.utils import get_column_letter

        # Styles
        title_font = Font(bold=True, size=14)
        section_font = Font(bold=True, size=11)
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1F4E78")
        subheader_fill = PatternFill("solid", fgColor="D9EAF7")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        def format_sheet(ws, df, sheet_title, start_row=6):
            # Sheet title
            ws["A1"] = sheet_title
            ws["A1"].font = title_font

            # Metadata block
            ws["A3"] = "Number of Setups"
            ws["B3"] = closure["number_of_setups"]
            ws["C3"] = "Total Distance"
            ws["D3"] = closure["total_distance"]
            ws["E3"] = "Precision Ratio"
            ws["F3"] = closure["precision_display"]

            for cell in ws[3]:
                cell.font = section_font
                cell.fill = subheader_fill
                cell.border = thin_border

            # Header row formatting
            for col_idx, col_name in enumerate(df.columns, start=1):
                cell = ws.cell(row=start_row, column=col_idx)
                cell.value = col_name
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")

            # Data rows formatting
            for row_idx in range(start_row + 1, start_row + 1 + len(df)):
                for col_idx in range(1, len(df.columns) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border

            # Freeze panes below headers
            ws.freeze_panes = f"A{start_row + 1}"

            # Auto-fit columns
            for col_idx, col_cells in enumerate(ws.columns, start=1):
                max_length = 0
                col_letter = get_column_letter(col_idx)
                for cell in col_cells:
                    try:
                        if cell.value is not None:
                            max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass
                ws.column_dimensions[col_letter].width = max_length + 2

        # Write dataframes starting at row 6
        raw_df.to_excel(writer, sheet_name="Raw Results", index=False, startrow=5)
        adjusted_df.to_excel(writer, sheet_name="Adjustment", index=False, startrow=5)
        coordinates_df.to_excel(writer, sheet_name="Adjusted Coordinates", index=False, startrow=5)

        ws_raw = writer.sheets["Raw Results"]
        ws_adj = writer.sheets["Adjustment"]
        ws_coords = writer.sheets["Adjusted Coordinates"]

        format_sheet(ws_raw, raw_df, "Raw Traverse Results")
        format_sheet(ws_adj, adjusted_df, "Compass Rule Adjustment")
        format_sheet(ws_coords, coordinates_df, "Adjusted Coordinates")

        # Number formatting for Raw Results
        raw_headers = {cell.value: cell.column for cell in ws_raw[6]}

        for col_name, fmt in {
            "Setup Number": "0",
            "Angle DD": "0.000000",
            "Azimuth": "0.000000",
            "Distance": "0.000",
            "Latitude": "0.000",
            "Departure": "0.000",
        }.items():
            if col_name in raw_headers:
                col = raw_headers[col_name]
                for row in range(7, 7 + len(raw_df)):
                    ws_raw.cell(row=row, column=col).number_format = fmt

        # Number formatting for Adjustment
        adj_headers = {cell.value: cell.column for cell in ws_adj[6]}
        for col_name in ["Setup Number"]:
            if col_name in adj_headers:
                col = adj_headers[col_name]
                for row in range(7, 7 + len(adjusted_df)):
                    ws_adj.cell(row=row, column=col).number_format = "0"

        for col_name in ["Distance", "Raw Latitude", "Adjusted Latitude", "Raw Departure", "Adjusted Departure"]:
            if col_name in adj_headers:
                col = adj_headers[col_name]
                for row in range(7, 7 + len(adjusted_df)):
                    ws_adj.cell(row=row, column=col).number_format = "0.000"

        for col_name in ["Lat Correction", "Dep Correction"]:
            if col_name in adj_headers:
                col = adj_headers[col_name]
                for row in range(7, 7 + len(adjusted_df)):
                    ws_adj.cell(row=row, column=col).number_format = "0.0000"

        # Number formatting for Coordinates
        coord_headers = {cell.value: cell.column for cell in ws_coords[6]}
        for col_name in ["northing", "easting"]:
            if col_name in coord_headers:
                col = coord_headers[col_name]
                for row in range(7, 7 + len(coordinates_df)):
                    ws_coords.cell(row=row, column=col).number_format = "0.000"

        # Traverse Plot Sheet
        chart_sheet = workbook.create_sheet(title="Traverse Plot")

        chart_sheet["A1"] = "Traverse Plot"
        chart_sheet["A1"].font = title_font

        chart_sheet["A3"] = "Point"
        chart_sheet["B3"] = "Easting"
        chart_sheet["C3"] = "Northing"

        for cell in chart_sheet[3]:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

        for i, row in enumerate(coordinates_df.itertuples(index=False), start=4):
            chart_sheet[f"A{i}"] = getattr(row, "point")
            chart_sheet[f"B{i}"] = float(getattr(row, "easting"))
            chart_sheet[f"C{i}"] = float(getattr(row, "northing"))

            chart_sheet[f"B{i}"].number_format = "0.000"
            chart_sheet[f"C{i}"].number_format = "0.000"

            chart_sheet[f"A{i}"].border = thin_border
            chart_sheet[f"B{i}"].border = thin_border
            chart_sheet[f"C{i}"].border = thin_border

        chart = ScatterChart()
        chart.title = "Adjusted Traverse Plot"
        chart.style = 2
        chart.x_axis.title = "Easting"
        chart.y_axis.title = "Northing"
        chart.height = 10
        chart.width = 18
        chart.legend = None

        xvalues = Reference(chart_sheet, min_col=2, min_row=4, max_row=len(coordinates_df) + 3)
        yvalues = Reference(chart_sheet, min_col=3, min_row=4, max_row=len(coordinates_df) + 3)

        series = Series(yvalues, xvalues, title="Traverse")
        series.smooth = False
        series.marker.symbol = "circle"
        series.marker.size = 7
        series.graphicalProperties.line.width = 19050

        chart.series.append(series)
        chart_sheet.add_chart(chart, "E3")

        chart_sheet.freeze_panes = "A4"

        # Auto-fit plot sheet columns
        for col_idx, col_cells in enumerate(chart_sheet.columns, start=1):
            max_length = 0
            col_letter = get_column_letter(col_idx)
            for cell in col_cells:
                try:
                    if cell.value is not None:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            chart_sheet.column_dimensions[col_letter].width = max_length + 2

    output.seek(0)
    return output.getvalue()




def build_save_data(settings, rows_df):
    clean_df = clean_first_setup_angles(rows_df.copy())
    rows = clean_df.where(pd.notnull(clean_df), None).to_dict(orient="records")

    return {
        "project_settings": {
            "traverse_name": settings.get("traverse_name", ""),
            "starting_point": settings.get("starting_point", 1),
            "starting_azimuth_deg": settings.get("starting_azimuth_deg", 0),
            "starting_azimuth_min": settings.get("starting_azimuth_min", 0),
            "starting_azimuth_sec": settings.get("starting_azimuth_sec", 0.0),
            "starting_northing": settings.get("starting_northing", 0.0),
            "starting_easting": settings.get("starting_easting", 0.0),
            "distance_units": settings.get("distance_units", "Feet"),
            "traverse_direction": settings.get("traverse_direction", "Clockwise"),
        },
        "num_setups": len(rows),
        "rows": rows,
    }


def save_data_to_json_bytes(save_data):
    return json.dumps(save_data, indent=2).encode("utf-8")


def load_save_file(uploaded_file):
    data = json.load(uploaded_file)

    project_settings = data.get("project_settings", {})
    rows = data.get("rows", [])
    num_setups = data.get("num_setups", len(rows) if rows else 4)

    rows_df = pd.DataFrame(rows)

    if rows_df.empty:
        rows_df = build_setup_rows(num_setups)

    return project_settings, rows_df, num_setups


st.set_page_config(page_title="Compass Rule Traverse Adjustment", layout="wide")
st.image("Banner.jpg", use_container_width=True)

st.title("Compass Rule Traverse Adjustment")
st.markdown(
    """
    Use this tool to enter traverse setups, calculate closure, apply the Compass Rule adjustment,
    and export results to CSV or Excel with a plotted traverse.
    """
)

if "project_settings" not in st.session_state:
    st.session_state.project_settings = SAMPLE_PROJECT.copy()

if "rows_df" not in st.session_state:
    st.session_state.rows_df = pd.DataFrame(SAMPLE_ROWS)

if "num_setups" not in st.session_state:
    st.session_state.num_setups = len(st.session_state.rows_df)

if "results" not in st.session_state:
    st.session_state.results = None

st.divider()

col1, col2 = st.columns([2, 1])

with col1:
    st.subheader("Project Setup")

    traverse_name = st.text_input(
        "Traverse Name",
        value=st.session_state.project_settings.get("traverse_name", "")
    )

    starting_point = st.text_input(
        "Starting Point",
        value=str(st.session_state.project_settings.get("starting_point", 1))
    )

    st.markdown("**Starting Azimuth from Point 1 to Point 2**")
    st.caption("Enter the starting azimuth in degrees, minutes, and seconds.")

    az_col1, az_col2, az_col3 = st.columns(3)

    with az_col1:
        starting_azimuth_deg = st.number_input(
            "Degrees",
            value=int(st.session_state.project_settings.get("starting_azimuth_deg", 0)),
            step=1
        )

    with az_col2:
        starting_azimuth_min = st.number_input(
            "Minutes",
            value=int(st.session_state.project_settings.get("starting_azimuth_min", 0)),
            step=1,
            min_value=0,
            max_value=59
        )

    with az_col3:
        starting_azimuth_sec = st.number_input(
            "Seconds",
            value=float(st.session_state.project_settings.get("starting_azimuth_sec", 0.0)),
            format="%.4f",
            min_value=0.0,
            max_value=59.9999
        )

    coord_col1, coord_col2 = st.columns(2)
    with coord_col1:
        starting_northing = st.number_input(
            "Starting Northing",
            value=float(st.session_state.project_settings.get("starting_northing", 0.0)),
            format="%.3f"
        )
    with coord_col2:
        starting_easting = st.number_input(
            "Starting Easting",
            value=float(st.session_state.project_settings.get("starting_easting", 0.0)),
            format="%.3f"
        )

    distance_units = st.selectbox(
        "Distance Units",
        options=["Feet", "Meters"],
        index=0 if st.session_state.project_settings.get("distance_units", "Feet") == "Feet" else 1
    )

    traverse_direction = st.radio(
        "Traverse Direction",
        options=["Clockwise", "Counterclockwise"],
        index=0 if st.session_state.project_settings.get("traverse_direction", "Clockwise") == "Clockwise" else 1,
        horizontal=True
    )

    st.caption(
        "Use Clockwise or Counterclockwise to match the turning direction of the traverse."
    )

with col2:
    st.subheader("Setup Builder")

    num_setups = st.number_input(
        "Number of Setups",
        min_value=2,
        max_value=100,
        value=int(st.session_state.num_setups),
        step=1
    )

    if st.button("Build Setup Table", use_container_width=True):
        st.session_state.num_setups = num_setups
        st.session_state.rows_df = build_setup_rows(num_setups)
        st.session_state.results = None
        st.rerun()

    st.subheader("Quick Actions")

    if st.button("Load Example", use_container_width=True):
        st.session_state.project_settings = SAMPLE_PROJECT.copy()
        st.session_state.rows_df = pd.DataFrame(SAMPLE_ROWS)
        st.session_state.num_setups = len(SAMPLE_ROWS)
        st.session_state.results = None
        st.rerun()

    if st.button("Reset", use_container_width=True):
        st.session_state.project_settings = {
            "traverse_name": "",
            "starting_point": 1,
            "starting_azimuth_deg": 0,
            "starting_azimuth_min": 0,
            "starting_azimuth_sec": 0.0,
            "starting_northing": 0.0,
            "starting_easting": 0.0,
            "distance_units": "Feet",
            "traverse_direction": "Clockwise",
        }
        st.session_state.num_setups = 4
        st.session_state.rows_df = build_setup_rows(4)
        st.session_state.results = None
        st.rerun()

st.divider()

st.subheader("Save / Load Traverse")
st.caption("Save the current traverse inputs to a JSON file, or load a previously saved traverse.")

save_settings = {
    "traverse_name": traverse_name,
    "starting_point": starting_point,
    "starting_azimuth_deg": starting_azimuth_deg,
    "starting_azimuth_min": starting_azimuth_min,
    "starting_azimuth_sec": starting_azimuth_sec,
    "starting_northing": starting_northing,
    "starting_easting": starting_easting,
    "distance_units": distance_units,
    "traverse_direction": traverse_direction,
}

save_data = build_save_data(save_settings, st.session_state.rows_df)
default_file_name = f"{(traverse_name or 'traverse').strip().replace(' ', '_')}.json"

save_col, load_col = st.columns(2)

with save_col:
    st.download_button(
        label="Save Traverse to JSON",
        data=save_data_to_json_bytes(save_data),
        file_name=default_file_name,
        mime="application/json",
        use_container_width=True
    )

with load_col:
    uploaded_file = st.file_uploader(
        "Load Traverse from JSON",
        type=["json"]
    )

    if uploaded_file is not None:
        try:
            loaded_settings, loaded_rows_df, loaded_num_setups = load_save_file(uploaded_file)
            st.session_state.project_settings = loaded_settings
            st.session_state.rows_df = loaded_rows_df
            st.session_state.num_setups = loaded_num_setups
            st.session_state.results = None
            st.success("Traverse loaded successfully. Review the data and click Calculate.")
        except Exception as e:
            st.error(f"Could not load save file: {e}")

st.divider()

st.subheader("Traverse Setups")
st.info("Setup 1 uses the starting azimuth and does not require interior angle values. Angle fields for Setup 1 are ignored.")

edited_df = st.data_editor(
    clean_first_setup_angles(st.session_state.rows_df.copy()),
    num_rows="dynamic",
    use_container_width=True,
    hide_index=True,
)

settings = {
    "traverse_name": traverse_name,
    "starting_point": starting_point,
    "starting_azimuth_deg": starting_azimuth_deg,
    "starting_azimuth_min": starting_azimuth_min,
    "starting_azimuth_sec": starting_azimuth_sec,
    "starting_northing": starting_northing,
    "starting_easting": starting_easting,
    "distance_units": distance_units,
    "traverse_direction": traverse_direction,
}

calc_col1, calc_col2 = st.columns([1, 3])
with calc_col1:
    calculate_clicked = st.button("Calculate", type="primary", use_container_width=True)

if calculate_clicked:
    edited_df = clean_first_setup_angles(edited_df.copy())

    st.session_state.project_settings = settings
    st.session_state.rows_df = edited_df
    st.session_state.num_setups = len(edited_df)

    rows = edited_df.to_dict(orient="records")

    project_errors = validate_project_settings(settings)
    row_errors, active_rows = validate_traverse_rows(rows)

    all_errors = project_errors + row_errors

    if all_errors:
        st.error("Please fix the following issues:")
        for err in all_errors:
            st.write(f"- {err}")
    else:
        try:
            results = calculate_traverse(settings, active_rows)
            st.session_state.results = results
            st.success("Calculation completed successfully.")
        except Exception as e:
            st.error(f"Calculation failed: {e}")

if st.session_state.results is not None:
    results = st.session_state.results
    closure = results["closure"]

    st.divider()
    st.subheader("Calculation Summary")

    summary_col1, summary_col2 = st.columns([2, 1])

    with summary_col1:
        st.markdown(
            f"""
            **Traverse Name:** {traverse_name or 'Untitled Traverse'}  
            **Starting Point:** {starting_point}  
            **Starting Azimuth (DMS):** {starting_azimuth_deg}° {starting_azimuth_min}' {starting_azimuth_sec}"  
            **Starting Azimuth (Decimal Degrees):** {results['starting_azimuth_decimal']:.6f}  
            **Traverse Direction:** {traverse_direction}  
            **Distance Units:** {distance_units}
            """
        )

    with summary_col2:
        st.metric("Number of Setups", closure["number_of_setups"])
        st.metric("Total Distance", f'{closure["total_distance"]:.3f}')
        st.metric("Precision Ratio", closure["precision_display"])

    metric_col1, metric_col2, metric_col3 = st.columns(3)
    metric_col1.metric("Sum of Latitude", f'{closure["sum_latitude"]:.6f}')
    metric_col2.metric("Sum of Departure", f'{closure["sum_departure"]:.6f}')
    metric_col3.metric("Linear Error of Closure", f'{closure["linear_error"]:.6f}')

    st.divider()
    st.subheader("Export Results")
    st.caption("Download the traverse outputs as Excel or CSV files.")

    excel_data = build_excel_workbook(results)

    export_col1, export_col2 = st.columns(2)
    with export_col1:
        st.download_button(
            label="Download Excel Workbook with Plot",
            data=excel_data,
            file_name="traverse_results_with_plot.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )

        st.download_button(
            label="Download Raw Results CSV",
            data=df_to_csv_download(results["raw_df"]),
            file_name="raw_traverse_results.csv",
            mime="text/csv",
            use_container_width=True
        )

    with export_col2:
        st.download_button(
            label="Download Adjustment Results CSV",
            data=df_to_csv_download(results["adjusted_df"]),
            file_name="compass_rule_adjustment.csv",
            mime="text/csv",
            use_container_width=True
        )

        st.download_button(
            label="Download Adjusted Coordinates CSV",
            data=df_to_csv_download(results["coordinates_df"]),
            file_name="adjusted_coordinates.csv",
            mime="text/csv",
            use_container_width=True
        )

    st.divider()
    st.subheader("Raw Traverse Results")
    st.dataframe(
        results["raw_df"].style.format({
            "Angle DD": "{:.6f}",
            "Azimuth": "{:.6f}",
            "Distance": "{:.3f}",
            "Latitude": "{:.3f}",
            "Departure": "{:.3f}",
        }, na_rep=""),
        use_container_width=True
    )

    st.subheader("Compass Rule Adjustment")
    st.dataframe(
        results["adjusted_df"].style.format({
            "Distance": "{:.3f}",
            "Raw Latitude": "{:.3f}",
            "Lat Correction": "{:.4f}",
            "Adjusted Latitude": "{:.3f}",
            "Raw Departure": "{:.3f}",
            "Dep Correction": "{:.4f}",
            "Adjusted Departure": "{:.3f}",
        }),
        use_container_width=True
    )

    st.subheader("Adjusted Coordinates")
    st.dataframe(
        results["coordinates_df"].style.format({
            "northing": "{:.3f}",
            "easting": "{:.3f}",
        }),
        use_container_width=True
    )
    st.subheader("Traverse Plot")
    plot_fig = build_traverse_plot(
        results["coordinates_df"],
        traverse_name=traverse_name or "Traverse Plot"
    )
    st.pyplot(plot_fig)








