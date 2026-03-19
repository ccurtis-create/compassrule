streamlit
pandas
numpy
openpyxl
matplotlib

SAMPLE_PROJECT = {
    "traverse_name": "Example Traverse",
    "starting_azimuth_deg": 0,
    "starting_azimuth_min": 0,
    "starting_azimuth_sec": 0.0,
    "starting_northing": 0.0,
    "starting_easting": 0.0,
    "distance_units": "Feet",
    "starting_point": 1,
    "traverse_direction": "Clockwise",
}

SAMPLE_ROWS = [
    {
        "setup_number": 1,
        "from_point": 1,
        "to_point": 2,
        "angle_deg": None,
        "angle_min": None,
        "angle_sec": None,
        "distance": 125.33,
    },
    {
        "setup_number": 2,
        "from_point": 2,
        "to_point": 3,
        "angle_deg": 95,
        "angle_min": 14,
        "angle_sec": 22.0,
        "distance": 240.17,
    },
    {
        "setup_number": 3,
        "from_point": 3,
        "to_point": 4,
        "angle_deg": 88,
        "angle_min": 20,
        "angle_sec": 14.0,
        "distance": 310.05,
    },
    {
        "setup_number": 4,
        "from_point": 4,
        "to_point": 1,
        "angle_deg": 102,
        "angle_min": 10,
        "angle_sec": 11.0,
        "distance": 198.42,
    },
]



def is_blank(value):
    return value is None or str(value).strip() == ""


def is_blank_row(row):
    fields = [
        "setup_number",
        "from_point",
        "to_point",
        "angle_deg",
        "angle_min",
        "angle_sec",
        "distance",
    ]
    return all(is_blank(row.get(field, "")) for field in fields)


def validate_project_settings(settings):
    errors = []

    starting_point = settings.get("starting_point", None)
    az_deg = settings.get("starting_azimuth_deg", None)
    az_min = settings.get("starting_azimuth_min", None)
    az_sec = settings.get("starting_azimuth_sec", None)

    if starting_point in [None, ""]:
        errors.append("Starting point is required.")

    for field_label, value in [
        ("Starting Azimuth Degrees", az_deg),
        ("Starting Azimuth Minutes", az_min),
        ("Starting Azimuth Seconds", az_sec),
    ]:
        if value in [None, ""]:
            errors.append(f"{field_label} is required.")

    try:
        if az_deg not in [None, ""]:
            float(az_deg)
    except ValueError:
        errors.append("Starting Azimuth Degrees must be numeric.")

    try:
        if az_min not in [None, ""]:
            az_min_num = float(az_min)
            if az_min_num < 0 or az_min_num > 59:
                errors.append("Starting Azimuth Minutes must be between 0 and 59.")
    except ValueError:
        errors.append("Starting Azimuth Minutes must be numeric.")

    try:
        if az_sec not in [None, ""]:
            az_sec_num = float(az_sec)
            if az_sec_num < 0 or az_sec_num >= 60:
                errors.append("Starting Azimuth Seconds must be between 0 and less than 60.")
    except ValueError:
        errors.append("Starting Azimuth Seconds must be numeric.")

    for field_name in ["starting_northing", "starting_easting"]:
        value = settings.get(field_name, 0)
        if value not in [None, ""]:
            try:
                float(value)
            except ValueError:
                errors.append(f"{field_name.replace('_', ' ').title()} must be numeric.")

    return errors


def validate_traverse_rows(rows):
    errors = []
    active_rows = []

    for idx, row in enumerate(rows, start=1):
        if is_blank_row(row):
            continue

        active_rows.append(row)

        setup_number = row.get("setup_number", "")
        from_point = row.get("from_point", "")
        to_point = row.get("to_point", "")
        angle_deg = row.get("angle_deg", "")
        angle_min = row.get("angle_min", "")
        angle_sec = row.get("angle_sec", "")
        distance = row.get("distance", "")

        if is_blank(setup_number):
            errors.append(f"Row {idx}: Setup Number is required.")

        if is_blank(from_point):
            errors.append(f"Row {idx}: From Point is required.")

        if is_blank(to_point):
            errors.append(f"Row {idx}: To Point is required.")

        if is_blank(distance):
            errors.append(f"Row {idx}: Distance is required.")

        try:
            if not is_blank(setup_number):
                float(setup_number)
        except ValueError:
            errors.append(f"Row {idx}: Setup Number must be numeric.")

        is_first_setup = False
        try:
            if not is_blank(setup_number):
                is_first_setup = float(setup_number) == 1
        except Exception:
            pass

        if not is_first_setup:
            for field_label, value in [
                ("Angle Degrees", angle_deg),
                ("Angle Minutes", angle_min),
                ("Angle Seconds", angle_sec),
            ]:
                if is_blank(value):
                    errors.append(f"Row {idx}: {field_label} is required for Setup 2 and higher.")

        try:
            if not is_blank(angle_deg):
                float(angle_deg)
        except ValueError:
            errors.append(f"Row {idx}: Angle Degrees must be numeric.")

        try:
            if not is_blank(angle_min):
                angle_min_num = float(angle_min)
                if angle_min_num < 0 or angle_min_num > 59:
                    errors.append(f"Row {idx}: Angle Minutes must be between 0 and 59.")
        except ValueError:
            if not is_first_setup:
                errors.append(f"Row {idx}: Angle Minutes must be numeric.")

        try:
            if not is_blank(angle_sec):
                angle_sec_num = float(angle_sec)
                if angle_sec_num < 0 or angle_sec_num >= 60:
                    errors.append(f"Row {idx}: Angle Seconds must be between 0 and less than 60.")
        except ValueError:
            if not is_first_setup:
                errors.append(f"Row {idx}: Angle Seconds must be numeric.")

        try:
            if not is_blank(distance):
                distance_num = float(distance)
                if distance_num <= 0:
                    errors.append(f"Row {idx}: Distance must be greater than zero.")
        except ValueError:
            errors.append(f"Row {idx}: Distance must be numeric.")

    if len(active_rows) < 2:
        errors.append("At least 2 active traverse setups are required.")

    return errors, active_rows


import io
import json
import matplotlib.pyplot as plt
import pandas as pd
import streamlit as st

from calculations import calculate_traverse
from sample_data import SAMPLE_PROJECT, SAMPLE_ROWS
from validation import validate_project_settings, validate_traverse_rows


def build_setup_rows(num_setups):
    rows = []
    for i in range(1, num_setups + 1):
        rows.append({
            "setup_number": i,
            "from_point": i,
            "to_point": i + 1 if i < num_setups else 1,
            "angle_deg": None,
            "angle_min": None,
            "angle_sec": None,
            "distance": None,
        })
    return pd.DataFrame(rows)


def clean_first_setup_angles(df):
    if not df.empty and "setup_number" in df.columns:
        for idx in df.index:
            try:
                if pd.notna(df.at[idx, "setup_number"]) and int(float(df.at[idx, "setup_number"])) == 1:
                    df.at[idx, "angle_deg"] = None
                    df.at[idx, "angle_min"] = None
                    df.at[idx, "angle_sec"] = None
            except Exception:
                pass
    return df


def df_to_csv_download(df):
    return df.to_csv(index=False).encode("utf-8")


def build_traverse_plot(coordinates_df, traverse_name="Traverse Plot"):
    fig, ax = plt.subplots(figsize=(10, 5))

    x = coordinates_df["easting"]
    y = coordinates_df["northing"]
    labels = coordinates_df["point"]

    ax.plot(x, y, marker="o", linestyle="-", linewidth=2, markersize=6)

    for xi, yi, label in zip(x, y, labels):
        ax.annotate(str(label), (xi, yi), textcoords="offset points", xytext=(5, 5))

    ax.set_title(traverse_name)
    ax.set_xlabel("Easting")
    ax.set_ylabel("Northing")
    ax.grid(True, linestyle="--", alpha=0.5)
    ax.set_aspect("equal", adjustable="datalim")

    plt.tight_layout()
    return fig

def build_excel_workbook(results):
    output = io.BytesIO()

    raw_df = results["raw_df"].copy()
    adjusted_df = results["adjusted_df"].copy()
    coordinates_df = results["coordinates_df"].copy()
    closure = results["closure"]

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        workbook = writer.book

        from openpyxl.chart import ScatterChart, Reference, Series
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.utils import get_column_letter

        # Styles
        title_font = Font(bold=True, size=14)
        section_font = Font(bold=True, size=11)
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1F4E78")
        subheader_fill = PatternFill("solid", fgColor="D9EAF7")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        def format_sheet(ws, df, sheet_title, start_row=6):
            # Sheet title
            ws["A1"] = sheet_title
            ws["A1"].font = title_font

            # Metadata block
            ws["A3"] = "Number of Setups"
            ws["B3"] = closure["number_of_setups"]
            ws["C3"] = "Total Distance"
            ws["D3"] = closure["total_distance"]
            ws["E3"] = "Precision Ratio"
            ws["F3"] = closure["precision_display"]

            for cell in ws[3]:
                cell.font = section_font
                cell.fill = subheader_fill
                cell.border = thin_border

            # Header row formatting
            for col_idx, col_name in enumerate(df.columns, start=1):
                cell = ws.cell(row=start_row, column=col_idx)
                cell.value = col_name
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")

            # Data rows formatting
            for row_idx in range(start_row + 1, start_row + 1 + len(df)):
                for col_idx in range(1, len(df.columns) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border

            # Freeze panes below headers
            ws.freeze_panes = f"A{start_row + 1}"

            # Auto-fit columns
            for col_idx, col_cells in enumerate(ws.columns, start=1):
                max_length = 0
                col_letter = get_column_letter(col_idx)
                for cell in col_cells:
                    try:
                        if cell.value is not None:
                            max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass
                ws.column_dimensions[col_letter].width = max_length + 2

        # Write dataframes starting at row 6
        raw_df.to_excel(writer, sheet_name="Raw Results", index=False, startrow=5)
        adjusted_df.to_excel(writer, sheet_name="Adjustment", index=False, startrow=5)
        coordinates_df.to_excel(writer, sheet_name="Adjusted Coordinates", index=False, startrow=5)

        ws_raw = writer.sheets["Raw Results"]
        ws_adj = writer.sheets["Adjustment"]
        ws_coords = writer.sheets["Adjusted Coordinates"]

        format_sheet(ws_raw, raw_df, "Raw Traverse Results")
        format_sheet(ws_adj, adjusted_df, "Compass Rule Adjustment")
        format_sheet(ws_coords, coordinates_df, "Adjusted Coordinates")

        # Number formatting for Raw Results
        raw_headers = {cell.value: cell.column for cell in ws_raw[6]}

        for col_name, fmt in {
            "Setup Number": "0",
            "Angle DD": "0.000000",
            "Azimuth": "0.000000",
            "Distance": "0.000",
            "Latitude": "0.000",
            "Departure": "0.000",
        }.items():
            if col_name in raw_headers:
                col = raw_headers[col_name]
                for row in range(7, 7 + len(raw_df)):
                    ws_raw.cell(row=row, column=col).number_format = fmt

        # Number formatting for Adjustment
        adj_headers = {cell.value: cell.column for cell in ws_adj[6]}
        for col_name in ["Setup Number"]:
            if col_name in adj_headers:
                col = adj_headers[col_name]
                for row in range(7, 7 + len(adjusted_df)):
                    ws_adj.cell(row=row, column=col).number_format = "0"

        for col_name in ["Distance", "Raw Latitude", "Adjusted Latitude", "Raw Departure", "Adjusted Departure"]:
            if col_name in adj_headers:
                col = adj_headers[col_name]
                for row in range(7, 7 + len(adjusted_df)):
                    ws_adj.cell(row=row, column=col).number_format = "0.000"

        for col_name in ["Lat Correction", "Dep Correction"]:
            if col_name in adj_headers:
                col = adj_headers[col_name]
                for row in range(7, 7 + len(adjusted_df)):
                    ws_adj.cell(row=row, column=col).number_format = "0.0000"

        # Number formatting for Coordinates
        coord_headers = {cell.value: cell.column for cell in ws_coords[6]}
        for col_name in ["northing", "easting"]:
            if col_name in coord_headers:
                col = coord_headers[col_name]
                for row in range(7, 7 + len(coordinates_df)):
                    ws_coords.cell(row=row, column=col).number_format = "0.000"

        # Traverse Plot Sheet
        chart_sheet = workbook.create_sheet(title="Traverse Plot")

        chart_sheet["A1"] = "Traverse Plot"
        chart_sheet["A1"].font = title_font

        chart_sheet["A3"] = "Point"
        chart_sheet["B3"] = "Easting"
        chart_sheet["C3"] = "Northing"

        for cell in chart_sheet[3]:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

        for i, row in enumerate(coordinates_df.itertuples(index=False), start=4):
            chart_sheet[f"A{i}"] = getattr(row, "point")
            chart_sheet[f"B{i}"] = float(getattr(row, "easting"))
            chart_sheet[f"C{i}"] = float(getattr(row, "northing"))

            chart_sheet[f"B{i}"].number_format = "0.000"
            chart_sheet[f"C{i}"].number_format = "0.000"

            chart_sheet[f"A{i}"].border = thin_border
            chart_sheet[f"B{i}"].border = thin_border
            chart_sheet[f"C{i}"].border = thin_border

        chart = ScatterChart()
        chart.title = "Adjusted Traverse Plot"
        chart.style = 2
        chart.x_axis.title = "Easting"
        chart.y_axis.title = "Northing"
        chart.height = 10
        chart.width = 18
        chart.legend = None

        xvalues = Reference(chart_sheet, min_col=2, min_row=4, max_row=len(coordinates_df) + 3)
        yvalues = Reference(chart_sheet, min_col=3, min_row=4, max_row=len(coordinates_df) + 3)

        series = Series(yvalues, xvalues, title="Traverse")
        series.smooth = False
        series.marker.symbol = "circle"
        series.marker.size = 7
        series.graphicalProperties.line.width = 19050

        chart.series.append(series)
        chart_sheet.add_chart(chart, "E3")

        chart_sheet.freeze_panes = "A4"

        # Auto-fit plot sheet columns
        for col_idx, col_cells in enumerate(chart_sheet.columns, start=1):
            max_length = 0
            col_letter = get_column_letter(col_idx)
            for cell in col_cells:
                try:
                    if cell.value is not None:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            chart_sheet.column_dimensions[col_letter].width = max_length + 2

    output.seek(0)
    return output.getvalue()




def build_save_data(settings, rows_df):
    clean_df = clean_first_setup_angles(rows_df.copy())
    rows = clean_df.where(pd.notnull(clean_df), None).to_dict(orient="records")

    return {
        "project_settings": {
            "traverse_name": settings.get("traverse_name", ""),
            "starting_point": settings.get("starting_point", 1),
            "starting_azimuth_deg": settings.get("starting_azimuth_deg", 0),
            "starting_azimuth_min": settings.get("starting_azimuth_min", 0),
            "starting_azimuth_sec": settings.get("starting_azimuth_sec", 0.0),
            "starting_northing": settings.get("starting_northing", 0.0),
            "starting_easting": settings.get("starting_easting", 0.0),
            "distance_units": settings.get("distance_units", "Feet"),
            "traverse_direction": settings.get("traverse_direction", "Clockwise"),
        },
        "num_setups": len(rows),
        "rows": rows,
    }


def save_data_to_json_bytes(save_data):
    return json.dumps(save_data, indent=2).encode("utf-8")


def load_save_file(uploaded_file):
    data = json.load(uploaded_file)

    project_settings = data.get("project_settings", {})
    rows = data.get("rows", [])
    num_setups = data.get("num_setups", len(rows) if rows else 4)

    rows_df = pd.DataFrame(rows)

    if rows_df.empty:
        rows_df = build_setup_rows(num_setups)

    return project_settings, rows_df, num_setups


st.set_page_config(page_title="Compass Rule Traverse Adjustment", layout="wide")
st.image("Banner.jpg", use_container_width=True)

st.title("Compass Rule Traverse Adjustment")
st.markdown(
    """
    Use this tool to enter traverse setups, calculate closure, apply the Compass Rule adjustment,
    and export results to CSV or Excel with a plotted traverse.
    """
)

if "project_settings" not in st.session_state:
    st.session_state.project_settings = SAMPLE_PROJECT.copy()

if "rows_df" not in st.session_state:
    st.session_state.rows_df = pd.DataFrame(SAMPLE_ROWS)

if "num_setups" not in st.session_state:
    st.session_state.num_setups = len(st.session_state.rows_df)

if "results" not in st.session_state:
    st.session_state.results = None

st.divider()

col1, col2 = st.columns([2, 1])

with col1:
    st.subheader("Project Setup")

    traverse_name = st.text_input(
        "Traverse Name",
        value=st.session_state.project_settings.get("traverse_name", "")
    )

    starting_point = st.text_input(
        "Starting Point",
        value=str(st.session_state.project_settings.get("starting_point", 1))
    )

    st.markdown("**Starting Azimuth from Point 1 to Point 2**")
    st.caption("Enter the starting azimuth in degrees, minutes, and seconds.")

    az_col1, az_col2, az_col3 = st.columns(3)

    with az_col1:
        starting_azimuth_deg = st.number_input(
            "Degrees",
            value=int(st.session_state.project_settings.get("starting_azimuth_deg", 0)),
            step=1
        )

    with az_col2:
        starting_azimuth_min = st.number_input(
            "Minutes",
            value=int(st.session_state.project_settings.get("starting_azimuth_min", 0)),
            step=1,
            min_value=0,
            max_value=59
        )

    with az_col3:
        starting_azimuth_sec = st.number_input(
            "Seconds",
            value=float(st.session_state.project_settings.get("starting_azimuth_sec", 0.0)),
            format="%.4f",
            min_value=0.0,
            max_value=59.9999
        )

    coord_col1, coord_col2 = st.columns(2)
    with coord_col1:
        starting_northing = st.number_input(
            "Starting Northing",
            value=float(st.session_state.project_settings.get("starting_northing", 0.0)),
            format="%.3f"
        )
    with coord_col2:
        starting_easting = st.number_input(
            "Starting Easting",
            value=float(st.session_state.project_settings.get("starting_easting", 0.0)),
            format="%.3f"
        )

    distance_units = st.selectbox(
        "Distance Units",
        options=["Feet", "Meters"],
        index=0 if st.session_state.project_settings.get("distance_units", "Feet") == "Feet" else 1
    )

    traverse_direction = st.radio(
        "Traverse Direction",
        options=["Clockwise", "Counterclockwise"],
        index=0 if st.session_state.project_settings.get("traverse_direction", "Clockwise") == "Clockwise" else 1,
        horizontal=True
    )

    st.caption(
        "Use Clockwise or Counterclockwise to match the turning direction of the traverse."
    )

with col2:
    st.subheader("Setup Builder")

    num_setups = st.number_input(
        "Number of Setups",
        min_value=2,
        max_value=100,
        value=int(st.session_state.num_setups),
        step=1
    )

    if st.button("Build Setup Table", use_container_width=True):
        st.session_state.num_setups = num_setups
        st.session_state.rows_df = build_setup_rows(num_setups)
        st.session_state.results = None
        st.rerun()

    st.subheader("Quick Actions")

    if st.button("Load Example", use_container_width=True):
        st.session_state.project_settings = SAMPLE_PROJECT.copy()
        st.session_state.rows_df = pd.DataFrame(SAMPLE_ROWS)
        st.session_state.num_setups = len(SAMPLE_ROWS)
        st.session_state.results = None
        st.rerun()

    if st.button("Reset", use_container_width=True):
        st.session_state.project_settings = {
            "traverse_name": "",
            "starting_point": 1,
            "starting_azimuth_deg": 0,
            "starting_azimuth_min": 0,
            "starting_azimuth_sec": 0.0,
            "starting_northing": 0.0,
            "starting_easting": 0.0,
            "distance_units": "Feet",
            "traverse_direction": "Clockwise",
        }
        st.session_state.num_setups = 4
        st.session_state.rows_df = build_setup_rows(4)
        st.session_state.results = None
        st.rerun()

st.divider()

st.subheader("Save / Load Traverse")
st.caption("Save the current traverse inputs to a JSON file, or load a previously saved traverse.")

save_settings = {
    "traverse_name": traverse_name,
    "starting_point": starting_point,
    "starting_azimuth_deg": starting_azimuth_deg,
    "starting_azimuth_min": starting_azimuth_min,
    "starting_azimuth_sec": starting_azimuth_sec,
    "starting_northing": starting_northing,
    "starting_easting": starting_easting,
    "distance_units": distance_units,
    "traverse_direction": traverse_direction,
}

save_data = build_save_data(save_settings, st.session_state.rows_df)
default_file_name = f"{(traverse_name or 'traverse').strip().replace(' ', '_')}.json"

save_col, load_col = st.columns(2)

with save_col:
    st.download_button(
        label="Save Traverse to JSON",
        data=save_data_to_json_bytes(save_data),
        file_name=default_file_name,
        mime="application/json",
        use_container_width=True
    )

with load_col:
    uploaded_file = st.file_uploader(
        "Load Traverse from JSON",
        type=["json"]
    )

    if uploaded_file is not None:
        try:
            loaded_settings, loaded_rows_df, loaded_num_setups = load_save_file(uploaded_file)
            st.session_state.project_settings = loaded_settings
            st.session_state.rows_df = loaded_rows_df
            st.session_state.num_setups = loaded_num_setups
            st.session_state.results = None
            st.success("Traverse loaded successfully. Review the data and click Calculate.")
        except Exception as e:
            st.error(f"Could not load save file: {e}")

st.divider()

st.subheader("Traverse Setups")
st.info("Setup 1 uses the starting azimuth and does not require interior angle values. Angle fields for Setup 1 are ignored.")

edited_df = st.data_editor(
    clean_first_setup_angles(st.session_state.rows_df.copy()),
    num_rows="dynamic",
    use_container_width=True,
    hide_index=True,
)

settings = {
    "traverse_name": traverse_name,
    "starting_point": starting_point,
    "starting_azimuth_deg": starting_azimuth_deg,
    "starting_azimuth_min": starting_azimuth_min,
    "starting_azimuth_sec": starting_azimuth_sec,
    "starting_northing": starting_northing,
    "starting_easting": starting_easting,
    "distance_units": distance_units,
    "traverse_direction": traverse_direction,
}

calc_col1, calc_col2 = st.columns([1, 3])
with calc_col1:
    calculate_clicked = st.button("Calculate", type="primary", use_container_width=True)

if calculate_clicked:
    edited_df = clean_first_setup_angles(edited_df.copy())

    st.session_state.project_settings = settings
    st.session_state.rows_df = edited_df
    st.session_state.num_setups = len(edited_df)

    rows = edited_df.to_dict(orient="records")

    project_errors = validate_project_settings(settings)
    row_errors, active_rows = validate_traverse_rows(rows)

    all_errors = project_errors + row_errors

    if all_errors:
        st.error("Please fix the following issues:")
        for err in all_errors:
            st.write(f"- {err}")
    else:
        try:
            results = calculate_traverse(settings, active_rows)
            st.session_state.results = results
            st.success("Calculation completed successfully.")
        except Exception as e:
            st.error(f"Calculation failed: {e}")

if st.session_state.results is not None:
    results = st.session_state.results
    closure = results["closure"]

    st.divider()
    st.subheader("Calculation Summary")

    summary_col1, summary_col2 = st.columns([2, 1])

    with summary_col1:
        st.markdown(
            f"""
            **Traverse Name:** {traverse_name or 'Untitled Traverse'}  
            **Starting Point:** {starting_point}  
            **Starting Azimuth (DMS):** {starting_azimuth_deg}° {starting_azimuth_min}' {starting_azimuth_sec}"  
            **Starting Azimuth (Decimal Degrees):** {results['starting_azimuth_decimal']:.6f}  
            **Traverse Direction:** {traverse_direction}  
            **Distance Units:** {distance_units}
            """
        )

    with summary_col2:
        st.metric("Number of Setups", closure["number_of_setups"])
        st.metric("Total Distance", f'{closure["total_distance"]:.3f}')
        st.metric("Precision Ratio", closure["precision_display"])

    metric_col1, metric_col2, metric_col3 = st.columns(3)
    metric_col1.metric("Sum of Latitude", f'{closure["sum_latitude"]:.6f}')
    metric_col2.metric("Sum of Departure", f'{closure["sum_departure"]:.6f}')
    metric_col3.metric("Linear Error of Closure", f'{closure["linear_error"]:.6f}')

    st.divider()
    st.subheader("Export Results")
    st.caption("Download the traverse outputs as Excel or CSV files.")

    excel_data = build_excel_workbook(results)

    export_col1, export_col2 = st.columns(2)
    with export_col1:
        st.download_button(
            label="Download Excel Workbook with Plot",
            data=excel_data,
            file_name="traverse_results_with_plot.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )

        st.download_button(
            label="Download Raw Results CSV",
            data=df_to_csv_download(results["raw_df"]),
            file_name="raw_traverse_results.csv",
            mime="text/csv",
            use_container_width=True
        )

    with export_col2:
        st.download_button(
            label="Download Adjustment Results CSV",
            data=df_to_csv_download(results["adjusted_df"]),
            file_name="compass_rule_adjustment.csv",
            mime="text/csv",
            use_container_width=True
        )

        st.download_button(
            label="Download Adjusted Coordinates CSV",
            data=df_to_csv_download(results["coordinates_df"]),
            file_name="adjusted_coordinates.csv",
            mime="text/csv",
            use_container_width=True
        )

    st.divider()
    st.subheader("Raw Traverse Results")
    st.dataframe(
        results["raw_df"].style.format({
            "Angle DD": "{:.6f}",
            "Azimuth": "{:.6f}",
            "Distance": "{:.3f}",
            "Latitude": "{:.3f}",
            "Departure": "{:.3f}",
        }, na_rep=""),
        use_container_width=True
    )

    st.subheader("Compass Rule Adjustment")
    st.dataframe(
        results["adjusted_df"].style.format({
            "Distance": "{:.3f}",
            "Raw Latitude": "{:.3f}",
            "Lat Correction": "{:.4f}",
            "Adjusted Latitude": "{:.3f}",
            "Raw Departure": "{:.3f}",
            "Dep Correction": "{:.4f}",
            "Adjusted Departure": "{:.3f}",
        }),
        use_container_width=True
    )

    st.subheader("Adjusted Coordinates")
    st.dataframe(
        results["coordinates_df"].style.format({
            "northing": "{:.3f}",
            "easting": "{:.3f}",
        }),
        use_container_width=True
    )
    st.subheader("Traverse Plot")
    plot_fig = build_traverse_plot(
        results["coordinates_df"],
        traverse_name=traverse_name or "Traverse Plot"
    )
    st.pyplot(plot_fig)









import math
import pandas as pd


def dms_to_decimal(deg, minute, sec):
    return float(deg) + float(minute) / 60.0 + float(sec) / 3600.0


def normalize_azimuth(angle):
    angle = float(angle) % 360.0
    if angle < 0:
        angle += 360.0
    return angle


def degrees_to_radians(angle_deg):
    return math.radians(angle_deg)


def build_clean_setups(rows):
    setups = []

    for row in rows:
        setup_number = int(float(row.get("setup_number", 0)))
        angle_deg = row.get("angle_deg", None)
        angle_min = row.get("angle_min", None)
        angle_sec = row.get("angle_sec", None)

        setup = {
            "setup_number": setup_number,
            "from_point": str(row.get("from_point", "")).strip(),
            "to_point": str(row.get("to_point", "")).strip(),
            "distance": float(row.get("distance", 0)),
            "angle_deg": None if angle_deg in [None, ""] else float(angle_deg),
            "angle_min": None if angle_min in [None, ""] else float(angle_min),
            "angle_sec": None if angle_sec in [None, ""] else float(angle_sec),
        }

        if setup_number == 1:
            setup["angle_decimal"] = None
        else:
            setup["angle_decimal"] = dms_to_decimal(
                setup["angle_deg"],
                setup["angle_min"],
                setup["angle_sec"],
            )

        setups.append(setup)

    setups.sort(key=lambda x: x["setup_number"])
    return setups


def compute_azimuths(starting_azimuth, setups, traverse_direction="Clockwise"):
    azimuths = []

    if not setups:
        return azimuths

    for i, setup in enumerate(setups):
        if i == 0:
            azimuths.append(normalize_azimuth(starting_azimuth))
        else:
            angle = setups[i]["angle_decimal"]

            if traverse_direction == "Counterclockwise":
                next_az = azimuths[i - 1] - 180.0 + angle
            else:
                next_az = azimuths[i - 1] + 180.0 - angle

            azimuths.append(normalize_azimuth(next_az))

    return azimuths


def compute_lat_dep(distance, azimuth_deg):
    az_rad = degrees_to_radians(azimuth_deg)
    latitude = float(distance) * math.cos(az_rad)
    departure = float(distance) * math.sin(az_rad)
    return latitude, departure


def compute_raw_results(setups, starting_azimuth, traverse_direction="Clockwise"):
    azimuths = compute_azimuths(starting_azimuth, setups, traverse_direction)

    for i, setup in enumerate(setups):
        setup["azimuth"] = azimuths[i]
        latitude, departure = compute_lat_dep(setup["distance"], setup["azimuth"])
        setup["raw_latitude"] = latitude
        setup["raw_departure"] = departure

    return setups


def compute_closure(setups):
    total_distance = sum(setup["distance"] for setup in setups)
    sum_latitude = sum(setup["raw_latitude"] for setup in setups)
    sum_departure = sum(setup["raw_departure"] for setup in setups)
    linear_error = math.sqrt(sum_latitude**2 + sum_departure**2)

    if linear_error > 0:
        precision_ratio = total_distance / linear_error
        precision_display = f"1:{round(precision_ratio):,}"
    else:
        precision_ratio = None
        precision_display = "Perfect Closure"

    return {
        "number_of_setups": len(setups),
        "total_distance": total_distance,
        "sum_latitude": sum_latitude,
        "sum_departure": sum_departure,
        "linear_error": linear_error,
        "precision_ratio": precision_ratio,
        "precision_display": precision_display,
    }


def apply_compass_rule(setups, closure):
    total_distance = closure["total_distance"]
    sum_latitude = closure["sum_latitude"]
    sum_departure = closure["sum_departure"]

    if total_distance <= 0:
        raise ValueError("Total distance must be greater than zero.")

    for setup in setups:
        lat_corr = -sum_latitude * (setup["distance"] / total_distance)
        dep_corr = -sum_departure * (setup["distance"] / total_distance)

        setup["latitude_correction"] = lat_corr
        setup["departure_correction"] = dep_corr
        setup["adjusted_latitude"] = setup["raw_latitude"] + lat_corr
        setup["adjusted_departure"] = setup["raw_departure"] + dep_corr

    return setups


def compute_adjusted_coordinates(setups, starting_northing=0.0, starting_easting=0.0, starting_point="1"):
    coordinates = []

    current_n = float(starting_northing)
    current_e = float(starting_easting)

    coordinates.append({
        "point": str(starting_point),
        "northing": current_n,
        "easting": current_e,
    })

    for setup in setups:
        current_n += setup["adjusted_latitude"]
        current_e += setup["adjusted_departure"]

        point_name = setup["to_point"] if setup["to_point"] else f"P{setup['setup_number'] + 1}"

        coordinates.append({
            "point": point_name,
            "northing": current_n,
            "easting": current_e,
        })

    return coordinates


def calculate_traverse(settings, rows):
    starting_azimuth = dms_to_decimal(
        settings["starting_azimuth_deg"],
        settings["starting_azimuth_min"],
        settings["starting_azimuth_sec"],
    )

    setups = build_clean_setups(rows)
    setups = compute_raw_results(
        setups,
        starting_azimuth,
        settings.get("traverse_direction", "Clockwise"),
    )

    closure = compute_closure(setups)
    setups = apply_compass_rule(setups, closure)
    coordinates = compute_adjusted_coordinates(
        setups,
        settings.get("starting_northing", 0.0),
        settings.get("starting_easting", 0.0),
        settings.get("starting_point", "1"),
    )

    raw_df = pd.DataFrame([
        {
            "Setup Number": s["setup_number"],
            "From Point": s["from_point"],
            "To Point": s["to_point"],
            "Angle DD": s["angle_decimal"],
            "Azimuth": s["azimuth"],
            "Distance": s["distance"],
            "Latitude": s["raw_latitude"],
            "Departure": s["raw_departure"],
        }
        for s in setups
    ])

    adjusted_df = pd.DataFrame([
        {
            "Setup Number": s["setup_number"],
            "Distance": s["distance"],
            "Raw Latitude": s["raw_latitude"],
            "Lat Correction": s["latitude_correction"],
            "Adjusted Latitude": s["adjusted_latitude"],
            "Raw Departure": s["raw_departure"],
            "Dep Correction": s["departure_correction"],
            "Adjusted Departure": s["adjusted_departure"],
        }
        for s in setups
    ])

    coordinates_df = pd.DataFrame(coordinates)

    return {
        "setups": setups,
        "closure": closure,
        "raw_df": raw_df,
        "adjusted_df": adjusted_df,
        "coordinates_df": coordinates_df,
        "starting_azimuth_decimal": starting_azimuth,
    }


# Compass Rule Traverse Adjustment App

This is a beginner-friendly Streamlit prototype for performing a Compass Rule traverse adjustment.

## Features
- Enter traverse lines
- Input internal angles in DMS
- Input distances
- Enter starting azimuth and starting coordinates
- Compute raw latitude and departure
- Compute closure
- Apply Compass Rule adjustment
- Display adjusted coordinates

## Setup

### 1. Create a folder
Create a folder named `compass-rule-app`

### 2. Add files
Add these files:
- `app.py`
- `calculations.py`
- `validation.py`
- `sample_data.py`
- `requirements.txt`

### 3. Install dependencies
Open a terminal in the project folder and run:

```bash
pip install -r requirements.txt
